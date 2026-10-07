"""Document ingestion: turn uploaded files into plain text.

Supports PDF, DOCX, HTML, TXT, CSV, and Excel. Scanned PDF pages and images
inside DOCX files are read with OCR (core/ocr.py). Output is always a Document
with plain text content — later phases (chunking, embedding) consume that,
not the original file.
"""

import io
import re
import zipfile
from dataclasses import dataclass, field

import pandas as pd
from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from docx.table import Table
from docx.text.paragraph import Paragraph
from PIL import Image
from pypdf import PdfReader

from core.logger import logger
from core.safety import MAX_DOCUMENT_CHARS, MAX_OCR_PAGES

# A PDF page with less extracted text than this is treated as scanned and OCR'd.
MIN_PAGE_TEXT_CHARS = 10
# Rendering scale for OCR (2x of 72 dpi, about 144 dpi).
OCR_RENDER_SCALE = 2
# Images smaller than this (px, either side) in a DOCX are icons/logos, not scans.
MIN_OCR_IMAGE_SIDE = 100


@dataclass
class Document:
    name: str
    text: str
    metadata: dict = field(default_factory=dict)


def _dataframe_to_text(df: pd.DataFrame) -> str:
    lines = []
    for _, row in df.iterrows():
        pairs = (f"{col}: {row[col]}" for col in df.columns)
        lines.append(", ".join(pairs))
    return "\n".join(lines)


def _ocr_page_images(pdf_bytes: bytes, page_numbers: list[int]) -> dict[int, str]:
    import pypdfium2 as pdfium

    from core.ocr import image_to_text

    pdf = pdfium.PdfDocument(pdf_bytes)
    try:
        return {
            n: image_to_text(pdf[n].render(scale=OCR_RENDER_SCALE).to_pil())
            for n in page_numbers
        }
    finally:
        pdf.close()


# PDF text cleanup (applies to pypdf-extracted pages only, never to OCR output).
_EDGE_LINES = 2  # running headers/footers sit within this many lines of a page's top/bottom
_MAX_HEADER_CHARS = 80
_TERMINAL_PUNCTUATION = (".", "!", "?", ":", ";")
_WRAP_FULLNESS = 0.8  # a line this close to the page's typical width is assumed hard-wrapped


def _repeat_key(line: str) -> str:
    # Digits are dropped so "File an appeal1" and "File an appeal2" (page number glued on) match.
    return re.sub(r"\d+", "", line).strip().lower()


def _drop_repeated_edge_lines(pages: list[list[str]]) -> list[list[str]]:
    """Drop running headers/footers: short lines at a page's edge that repeat on other pages.

    The first occurrence is kept (it still labels the section); later ones are dropped.
    Needs at least 4 pages, so short documents are never touched.
    """
    if len(pages) < 4:
        return pages

    def edge_indexes(lines):
        idx = [i for i, line in enumerate(lines) if line]
        return set(idx[:_EDGE_LINES] + idx[-_EDGE_LINES:])

    page_count = {}
    for lines in pages:
        keys = {_repeat_key(lines[i]) for i in edge_indexes(lines) if len(lines[i]) <= _MAX_HEADER_CHARS}
        for key in keys:
            page_count[key] = page_count.get(key, 0) + 1

    seen = set()
    cleaned = []
    for lines in pages:
        edges = edge_indexes(lines)
        kept = []
        for i, line in enumerate(lines):
            key = _repeat_key(line)
            if i in edges and len(line) <= _MAX_HEADER_CHARS and page_count.get(key, 0) >= 2:
                if not key or key in seen:  # empty key = bare page number
                    continue
                seen.add(key)
            kept.append(line)
        cleaned.append(kept)
    return cleaned


def _reflow(lines: list[str], wrap_width: int) -> list[str]:
    """Join hard-wrapped lines back into paragraphs; headings and list items stay separate."""
    out = []
    for line in lines:
        if not line:
            out.append(line)
            continue
        prev = out[-1] if out else ""
        wrapped = (
            prev
            and not prev.endswith(_TERMINAL_PUNCTUATION)
            and not line.startswith(("- ", "#"))
            and (line[0].islower() or len(prev) >= _WRAP_FULLNESS * wrap_width)
        )
        if wrapped:
            out[-1] = prev + ("" if prev.endswith("-") else " ") + line
        else:
            out.append(line)
    return out


def _clean_pdf_pages(page_texts: list[str], skip: set[int]) -> list[str]:
    """Tidy pypdf output. Pages in `skip` (OCR'd) are returned unchanged."""
    split = []
    for i, text in enumerate(page_texts):
        if i in skip:
            split.append([])
            continue
        text = text.replace("•", "-").replace(" ", " ")
        lines = (re.sub(r"[ \t]+", " ", line).strip() for line in text.splitlines())
        split.append([re.sub(r" ([,.;:!?])", r"\1", line) for line in lines])  # "MSN ." -> "MSN."

    split = _drop_repeated_edge_lines([p if i not in skip else [] for i, p in enumerate(split)])

    lengths = sorted(len(line) for lines in split for line in lines if line)
    wrap_width = lengths[int(0.9 * (len(lengths) - 1))] if lengths else 0

    result = []
    for i, text in enumerate(page_texts):
        if i in skip:
            result.append(text)
            continue
        paragraph_lines = "\n".join(_reflow(split[i], wrap_width))
        result.append(re.sub(r"\n{3,}", "\n\n", paragraph_lines).strip())
    return result


def parse_pdf(file) -> str:
    reader = PdfReader(file)
    pages = [(page.extract_text() or "") for page in reader.pages]

    scanned = [i for i, text in enumerate(pages) if len(text.strip()) < MIN_PAGE_TEXT_CHARS]
    if scanned:
        if len(scanned) > MAX_OCR_PAGES:
            raise ValueError(
                f"{len(scanned)} pages have no text and would need OCR; the limit is {MAX_OCR_PAGES}"
            )
        logger.info("OCR on %d PDF page(s) with no text layer", len(scanned))
        file.seek(0)
        for i, text in _ocr_page_images(file.read(), scanned).items():
            pages[i] = text

    return "\n".join(_clean_pdf_pages(pages, skip=set(scanned)))


def _clean_inline(text: str) -> str:
    text = re.sub(r"\s+", " ", text).strip()
    return re.sub(r"\s+([,.;:!?])", r"\1", text)  # "Medicare ," -> "Medicare,"


def parse_html(file) -> str:
    soup = BeautifulSoup(file.read(), "lxml")
    # Page chrome and non-content: menus, footers, signup forms, scripts.
    for tag in soup(["script", "style", "noscript", "svg", "nav", "footer", "header", "aside", "form", "button", "input"]):
        tag.decompose()

    root = soup.find("main") or soup.body or soup
    block_tags = ["h1", "h2", "h3", "h4", "p", "li", "table"]
    lines = []
    for el in root.find_all(block_tags):
        if el.find_parent(block_tags):  # already covered by its parent (e.g. <p> inside <li>)
            continue
        text = _clean_inline(el.get_text(" ", strip=True))
        if not text:
            continue
        if el.name[0] == "h":
            lines.append(f"\n{'#' * int(el.name[1])} {text}")
        elif el.name == "li":
            lines.append(f"- {text}")
        else:
            lines.append(text)
    return "\n".join(lines).strip()


def _docx_table_to_text(table: Table) -> str:
    rows = [[_clean_inline(cell.text) for cell in row.cells] for row in table.rows]
    if not rows:
        return ""
    header, body = rows[0], rows[1:]
    if not body:
        return ", ".join(c for c in header if c)
    # Same "col: val, col: val" shape as CSV rows, so each row stays self-describing.
    return "\n".join(
        ", ".join(f"{h}: {v}" if h else v for h, v in zip(header, row) if v) for row in body
    )


def _docx_image_text(file) -> str:
    from core.ocr import image_to_text

    file.seek(0)
    with zipfile.ZipFile(file) as archive:
        names = sorted(n for n in archive.namelist() if n.startswith("word/media/"))
        images = []
        for name in names:
            try:
                image = Image.open(io.BytesIO(archive.read(name)))
                image.load()
            except Exception:
                continue  # unsupported image format (e.g. EMF/WMF)
            if min(image.size) >= MIN_OCR_IMAGE_SIDE:
                images.append(image)

    if len(images) > MAX_OCR_PAGES:
        logger.warning("DOCX has %d images; OCR'ing only the first %d", len(images), MAX_OCR_PAGES)
        images = images[:MAX_OCR_PAGES]
    if images:
        logger.info("OCR on %d image(s) embedded in a DOCX", len(images))
    return "\n".join(t for t in (image_to_text(img) for img in images) if t.strip())


def parse_docx(file) -> str:
    doc = DocxDocument(file)
    parts = []
    for child in doc.element.body.iterchildren():
        if child.tag.endswith("}p"):
            paragraph = Paragraph(child, doc)
            text = paragraph.text.strip()
            if not text:
                continue
            heading = re.match(r"Heading (\d)", paragraph.style.name if paragraph.style is not None else "")
            parts.append(f"\n{'#' * int(heading.group(1))} {text}" if heading else text)
        elif child.tag.endswith("}tbl"):
            parts.append(_docx_table_to_text(Table(child, doc)))

    image_text = _docx_image_text(file)
    if image_text:
        parts.append(image_text)
    return "\n".join(parts).strip()


def parse_txt(file) -> str:
    return file.read().decode("utf-8", errors="replace")


def parse_csv(file) -> str:
    df = pd.read_csv(file)
    return _dataframe_to_text(df)


def parse_excel(file) -> str:
    sheets = pd.read_excel(file, sheet_name=None, engine="openpyxl")
    parts = []
    for sheet_name, df in sheets.items():
        parts.append(f"[Sheet: {sheet_name}]\n{_dataframe_to_text(df)}")
    return "\n\n".join(parts)


_PARSERS = {
    "pdf": parse_pdf,
    "docx": parse_docx,
    "html": parse_html,
    "htm": parse_html,
    "txt": parse_txt,
    "csv": parse_csv,
    "xlsx": parse_excel,
    "xls": parse_excel,
}


def ingest(file) -> Document:
    extension = file.name.rsplit(".", 1)[-1].lower()
    parser = _PARSERS.get(extension)

    if parser is None:
        return Document(name=file.name, text="", metadata={"error": f"Unsupported file type: .{extension}"})

    file.seek(0)
    try:
        text = parser(file)
    except Exception as e:
        return Document(name=file.name, text="", metadata={"error": str(e)})

    if len(text) > MAX_DOCUMENT_CHARS:
        return Document(
            name=file.name,
            text="",
            metadata={"error": f"extracted text is too large ({len(text):,} characters; limit {MAX_DOCUMENT_CHARS:,})"},
        )

    return Document(name=file.name, text=text)
